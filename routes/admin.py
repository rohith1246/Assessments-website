from app import csrf
"""
Lionix — Admin Blueprint
Handles: login/logout, dashboard, assessment CRUD,
         question CRUD, results viewing, export, and record deletion/clean-up.
"""
from datetime import datetime
from flask import Blueprint, render_template, redirect, url_for, flash, request, session, jsonify, make_response, send_file
from flask_login import login_user, logout_user, login_required, current_user
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload

from models.models import db, Admin, Assessment, Question, Submission, Candidate, Answer, CodingSubmission, CodingProblem
from services.stats_service import get_dashboard_stats, get_recent_coding_submissions
from extensions import cache
from services.export_service import (
    export_csv, export_xlsx, export_daily_reports_csv, export_daily_reports_xlsx, _get_daily_report_data
)
from utils.helpers import sanitize_string, api_success, api_error, paginate_query
from sqlalchemy import func

admin_bp = Blueprint('admin', __name__, url_prefix='/admin')


# ─────────────────────────────────────────────
# AUTH
# ─────────────────────────────────────────────

@admin_bp.route('/login', methods=['GET', 'POST'])
@csrf.exempt
def login():
    if current_user.is_authenticated:
        return redirect(url_for('admin.dashboard'))

    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')

        if not email or not password:
            flash('Email and password are required.', 'danger')
            return render_template('admin/login.html')

        try:
            admin = Admin.query.filter_by(email=email).first()
            if admin and admin.check_password(password):
                login_user(admin, remember=False)
                next_page = request.args.get('next')
                flash(f'Welcome back, {admin.email}!', 'success')
                return redirect(next_page or url_for('admin.dashboard'))
            else:
                flash('Invalid email or password.', 'danger')
        except Exception as err:
            db.session.rollback()
            flash(f'Login error: {str(err)}', 'danger')

    return render_template('admin/login.html')


@admin_bp.route('/logout')
@login_required
def logout():
    logout_user()
    flash('You have been logged out.', 'info')
    return redirect(url_for('admin.login'))


# ─────────────────────────────────────────────
# DASHBOARD
# ─────────────────────────────────────────────

@admin_bp.route('/')
@admin_bp.route('/dashboard')
@login_required
def dashboard():
    try:
        stats = get_dashboard_stats()
        recent_results = (
            db.session.query(Submission)
            .options(joinedload(Submission.candidate), joinedload(Submission.assessment))
            .filter(Submission.status.in_(['pass', 'fail']))
            .order_by(Submission.submitted_at.desc())
            .limit(20)
            .all()
        )
        passed_results = [r for r in recent_results if r.status == 'pass'][:10]
        failed_results = [r for r in recent_results if r.status == 'fail'][:10]
        assessments = Assessment.query.order_by(Assessment.created_at.desc()).all()
        recent_coding_submissions = get_recent_coding_submissions(limit=15)

        return render_template(
            'admin/dashboard.html',
            stats=stats,
            assessments=assessments,
            passed_results=passed_results,
            failed_results=failed_results,
        )
    except Exception as e:
        db.session.rollback()
        return render_template(
            'admin/dashboard.html',
            stats={'total_candidates': 0, 'total_assessments': 0, 'active_assessments': 0, 'total_attempts': 0, 'passed': 0, 'failed': 0, 'in_progress': 0, 'pass_rate': 0},
            passed_results=[],
            failed_results=[],
            assessments=[],
            all_assessments=[]
        )


# ─────────────────────────────────────────────
# ASSESSMENT MANAGEMENT
# ─────────────────────────────────────────────

@admin_bp.route('/assessments')
@login_required
def assessments():
    # Annotate question_count via subquery to avoid N+1
    q_count_subq = (
        db.session.query(
            Question.assessment_id,
            func.count(Question.id).label('q_count')
        )
        .group_by(Question.assessment_id)
        .subquery()
    )
    all_assessments = (
        db.session.query(Assessment, q_count_subq.c.q_count)
        .outerjoin(q_count_subq, Assessment.id == q_count_subq.c.assessment_id)
        .order_by(Assessment.created_at.desc())
        .all()
    )
    # Attach the count so templates/to_dict can access it
    for assessment, q_count in all_assessments:
        assessment._question_count = q_count or 0
    assessment_list = [a for a, _ in all_assessments]
    return render_template('admin/assessments.html', assessments=assessment_list)


@admin_bp.route('/assessments/create', methods=['POST'])
@login_required
def create_assessment():
    title = sanitize_string(request.form.get('title', ''))
    description = sanitize_string(request.form.get('description', ''), max_length=1000)
    duration = request.form.get('duration', '60')
    pass_pct = request.form.get('pass_percentage', '60')

    if not title:
        flash('Assessment title is required.', 'danger')
        return redirect(url_for('admin.assessments'))

    try:
        duration = int(duration)
        pass_pct = float(pass_pct)
        if duration < 1 or pass_pct < 0 or pass_pct > 100:
            raise ValueError
    except ValueError:
        flash('Invalid duration or pass percentage.', 'danger')
        return redirect(url_for('admin.assessments'))

    assessment = Assessment(
        title=title,
        description=description,
        duration=duration,
        pass_percentage=pass_pct,
        status='inactive'
    )
    db.session.add(assessment)
    db.session.commit()
    flash(f'Assessment "{title}" created successfully.', 'success')
    return redirect(url_for('admin.assessments'))


@admin_bp.route('/assessments/<int:assessment_id>/edit', methods=['POST'])
@login_required
def edit_assessment(assessment_id):
    assessment = Assessment.query.get_or_404(assessment_id)
    assessment.title = sanitize_string(request.form.get('title', assessment.title))
    assessment.description = sanitize_string(
        request.form.get('description', assessment.description or ''), max_length=1000
    )
    try:
        assessment.duration = int(request.form.get('duration', assessment.duration))
        assessment.pass_percentage = float(
            request.form.get('pass_percentage', assessment.pass_percentage)
        )
    except ValueError:
        flash('Invalid duration or pass percentage.', 'danger')
        return redirect(url_for('admin.assessments'))

    db.session.commit()
    flash('Assessment updated successfully.', 'success')
    return redirect(url_for('admin.assessments'))


@admin_bp.route('/assessments/<int:assessment_id>/delete', methods=['POST'])
@login_required
def delete_assessment(assessment_id):
    assessment = Assessment.query.get_or_404(assessment_id)
    title = assessment.title
    db.session.delete(assessment)
    db.session.commit()
    flash(f'Assessment "{title}" deleted.', 'success')
    return redirect(url_for('admin.assessments'))


@admin_bp.route('/assessments/<int:assessment_id>/toggle-status', methods=['POST'])
@login_required
def toggle_assessment_status(assessment_id):
    assessment = Assessment.query.get_or_404(assessment_id)

    if assessment.status == 'inactive':
        assessment.status = 'active'
        msg = f'Assessment "{assessment.title}" is now ACTIVE.'
    else:
        assessment.status = 'inactive'
        msg = f'Assessment "{assessment.title}" deactivated.'

    db.session.commit()
    flash(msg, 'success')
    return redirect(url_for('admin.assessments'))


# ─────────────────────────────────────────────
# QUESTION BANK MANAGEMENT
# ─────────────────────────────────────────────

@admin_bp.route('/assessments/<int:assessment_id>/questions')
@login_required
def questions(assessment_id):
    assessment = Assessment.query.get_or_404(assessment_id)
    page = request.args.get('page', 1, type=int)
    per_page = 20
    q_query = Question.query.filter_by(assessment_id=assessment_id).order_by(Question.id)
    pagination = q_query.paginate(page=page, per_page=per_page, error_out=False)
    return render_template(
        'admin/questions.html',
        assessment=assessment,
        questions=pagination.items,
        pagination=pagination,
    )


@admin_bp.route('/assessments/<int:assessment_id>/questions/add', methods=['POST'])
@login_required
def add_question(assessment_id):
    assessment = Assessment.query.get_or_404(assessment_id)
    question_text = request.form.get('question', '').strip()
    option_a = request.form.get('option_a', '').strip()
    option_b = request.form.get('option_b', '').strip()
    option_c = request.form.get('option_c', '').strip()
    option_d = request.form.get('option_d', '').strip()
    correct_answer = request.form.get('correct_answer', '').strip().upper()

    if not all([question_text, option_a, option_b, option_c, option_d, correct_answer]):
        flash('All fields are required.', 'danger')
        return redirect(url_for('admin.questions', assessment_id=assessment_id))

    if correct_answer not in ('A', 'B', 'C', 'D'):
        flash('Correct answer must be A, B, C, or D.', 'danger')
        return redirect(url_for('admin.questions', assessment_id=assessment_id))

    q = Question(
        assessment_id=assessment_id,
        question=question_text,
        option_a=option_a,
        option_b=option_b,
        option_c=option_c,
        option_d=option_d,
        correct_answer=correct_answer,
    )
    db.session.add(q)
    db.session.commit()
    flash('Question added successfully.', 'success')
    return redirect(url_for('admin.questions', assessment_id=assessment_id))


@admin_bp.route('/questions/<int:question_id>/edit', methods=['POST'])
@login_required
def edit_question(question_id):
    q = Question.query.get_or_404(question_id)
    q.question = request.form.get('question', q.question).strip()
    q.option_a = request.form.get('option_a', q.option_a).strip()
    q.option_b = request.form.get('option_b', q.option_b).strip()
    q.option_c = request.form.get('option_c', q.option_c).strip()
    q.option_d = request.form.get('option_d', q.option_d).strip()
    correct = request.form.get('correct_answer', q.correct_answer).strip().upper()
    if correct in ('A', 'B', 'C', 'D'):
        q.correct_answer = correct
    db.session.commit()
    flash('Question updated.', 'success')
    return redirect(url_for('admin.questions', assessment_id=q.assessment_id))


@admin_bp.route('/questions/<int:question_id>/delete', methods=['POST'])
@login_required
def delete_question(question_id):
    q = Question.query.get_or_404(question_id)
    assessment_id = q.assessment_id
    db.session.delete(q)
    db.session.commit()
    flash('Question deleted.', 'success')
    return redirect(url_for('admin.questions', assessment_id=assessment_id))


# ─────────────────────────────────────────────
# RESULTS
# ─────────────────────────────────────────────

@admin_bp.route('/results')
@login_required
def results():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '').strip()
    status = request.args.get('status', 'all').strip().lower()
    per_page = 20

    query = (
        db.session.query(Submission)
        .options(
            joinedload(Submission.candidate),
            joinedload(Submission.assessment)
        )
        .filter(Submission.status != 'in_progress')
    )

    if status in ('pass', 'fail'):
        query = query.filter(Submission.status == status)

    query = query.order_by(
        Submission.percentage.desc(),
        Submission.score.desc(),
        Submission.submitted_at.desc()
    )

    if search:
        from models.models import Candidate as Cand
        like = f'%{search}%'
        query = query.join(Submission.candidate).filter(
            db.or_(
                Cand.full_name.ilike(like),
                Cand.hall_ticket.ilike(like),
                Cand.email.ilike(like),
            )
        )

    total = query.count()
    submissions = query.offset((page - 1) * per_page).limit(per_page).all()
    total_pages = max(1, (total + per_page - 1) // per_page)
    all_assessments = Assessment.query.order_by(Assessment.title).all()

    return render_template(
        'admin/results.html',
        submissions=submissions,
        page=page,
        total_pages=total_pages,
        total=total,
        search=search,
        status=status,
        per_page=per_page,
        all_assessments=all_assessments,
    )


@admin_bp.route('/results/export')
@login_required
def export_results():
    fmt = request.args.get('fmt', 'csv').lower()
    search = request.args.get('search', '').strip()
    status = request.args.get('status', None)
    assessment_id = request.args.get('assessment_id', None, type=int)

    if fmt == 'xlsx':
        return export_xlsx(search=search, assessment_id=assessment_id, status=status)
    return export_csv(search=search, assessment_id=assessment_id, status=status)


# ─────────────────────────────────────────────
# SUBMISSION / RECORD DELETION
# ─────────────────────────────────────────────

@admin_bp.route('/submissions/<int:submission_id>/delete', methods=['POST'])
@login_required
def delete_submission(submission_id):
    submission = Submission.query.get_or_404(submission_id)
    cand_name = submission.candidate.full_name if submission.candidate else f"Candidate #{submission.candidate_id}"
    candidate_id = submission.candidate_id
    
    # Delete answers & coding submissions for this submission first
    Answer.query.filter_by(submission_id=submission.id).delete(synchronize_session=False)
    CodingSubmission.query.filter_by(submission_id=submission.id).delete(synchronize_session=False)
    db.session.delete(submission)
    db.session.commit()

    # Check if candidate has any remaining submissions; if not, delete candidate record too
    remaining = Submission.query.filter_by(candidate_id=candidate_id).count()
    if remaining == 0:
        Candidate.query.filter_by(id=candidate_id).delete(synchronize_session=False)
        db.session.commit()
    
    cache.clear()
    flash(f'Assessment record for {cand_name} (ID #{submission_id}) has been deleted.', 'success')
    return redirect(request.referrer or url_for('admin.results'))


@admin_bp.route('/submissions/bulk-delete', methods=['POST'])
@login_required
def bulk_delete_submissions():
    submission_ids = request.form.getlist('submission_ids')
    delete_scope = request.form.get('delete_scope', 'selected')
    assessment_id = request.form.get('assessment_id', type=int)
    before_date_str = request.form.get('before_date', '').strip()
    status_filter = request.form.get('status_filter', '').strip().lower()
    wipe_all = request.form.get('wipe_all_candidates') == '1' or delete_scope == 'reset_drive'

    # Option 1: Complete Portal & Recruitment Drive Reset
    if wipe_all:
        Answer.query.delete(synchronize_session=False)
        CodingSubmission.query.delete(synchronize_session=False)
        Submission.query.delete(synchronize_session=False)
        Candidate.query.delete(synchronize_session=False)
        db.session.commit()
        cache.clear()
        flash('Successfully reset the entire campus drive. All candidate registrations, test attempts, answers, and coding submissions have been wiped clean.', 'success')
        return redirect(request.referrer or url_for('admin.dashboard'))

    # Option 2: Selected Checkbox Deletion
    if submission_ids and delete_scope == 'selected':
        sub_ids = [int(i) for i in submission_ids if str(i).isdigit()]
        if sub_ids:
            # Find associated candidate IDs before deleting
            cand_ids = [s.candidate_id for s in Submission.query.filter(Submission.id.in_(sub_ids)).all()]
            
            Answer.query.filter(Answer.submission_id.in_(sub_ids)).delete(synchronize_session=False)
            CodingSubmission.query.filter(CodingSubmission.submission_id.in_(sub_ids)).delete(synchronize_session=False)
            deleted_count = Submission.query.filter(Submission.id.in_(sub_ids)).delete(synchronize_session=False)
            db.session.commit()

            # Clean orphaned candidates
            if cand_ids:
                for cid in set(cand_ids):
                    if Submission.query.filter_by(candidate_id=cid).count() == 0:
                        Candidate.query.filter_by(id=cid).delete(synchronize_session=False)
                db.session.commit()

            cache.clear()
            flash(f'Successfully deleted {deleted_count} selected assessment record(s).', 'success')
            return redirect(request.referrer or url_for('admin.results'))
        else:
            flash('No valid records selected for deletion.', 'warning')
            return redirect(request.referrer or url_for('admin.results'))

    # Option 3: Filtered Deletion
    query = Submission.query

    if assessment_id:
        query = query.filter(Submission.assessment_id == assessment_id)

    if before_date_str:
        try:
            before_date = datetime.strptime(before_date_str, '%Y-%m-%d')
            query = query.filter(Submission.submitted_at < before_date)
        except ValueError:
            flash('Invalid date format provided for clean-up.', 'danger')
            return redirect(request.referrer or url_for('admin.results'))

    if status_filter in ('pass', 'fail', 'in_progress'):
        query = query.filter(Submission.status == status_filter)

    records = query.all()
    count = len(records)

    if count == 0:
        flash('No assessment records matched the specified criteria.', 'info')
        return redirect(request.referrer or url_for('admin.results'))

    rec_ids = [r.id for r in records]
    cand_ids = [r.candidate_id for r in records]

    Answer.query.filter(Answer.submission_id.in_(rec_ids)).delete(synchronize_session=False)
    CodingSubmission.query.filter(CodingSubmission.submission_id.in_(rec_ids)).delete(synchronize_session=False)
    Submission.query.filter(Submission.id.in_(rec_ids)).delete(synchronize_session=False)
    db.session.commit()

    # Clean orphaned candidates
    if cand_ids:
        for cid in set(cand_ids):
            if Submission.query.filter_by(candidate_id=cid).count() == 0:
                Candidate.query.filter_by(id=cid).delete(synchronize_session=False)
        db.session.commit()

    cache.clear()
    flash(f'Successfully deleted {count} old assessment record(s).', 'success')
    return redirect(request.referrer or url_for('admin.results'))


# ─────────────────────────────────────────────
# API — Coding Submission Source Code for Modal
# ─────────────────────────────────────────────

@admin_bp.route('/api/coding-submissions/<int:coding_sub_id>')
@login_required
def api_coding_submission(coding_sub_id):
    cs = CodingSubmission.query.get_or_404(coding_sub_id)
    return jsonify({
        'id': cs.id,
        'candidate_name': cs.submission.candidate.full_name if cs.submission and cs.submission.candidate else 'Unknown',
        'hall_ticket': cs.submission.candidate.hall_ticket if cs.submission and cs.submission.candidate else 'N/A',
        'email': cs.submission.candidate.email if cs.submission and cs.submission.candidate else 'N/A',
        'problem_title': cs.problem.title if cs.problem else 'Coding Challenge',
        'language': cs.language,
        'source_code': cs.source_code,
        'passed_testcases': cs.passed_testcases,
        'total_testcases': cs.total_testcases,
        'score': cs.score,
        'execution_time_ms': cs.execution_time_ms,
        'status': cs.status,
        'submitted_at': cs.submitted_at.strftime('%d %b %Y, %I:%M %p') if cs.submitted_at else 'N/A'
    })


# ─────────────────────────────────────────────
# API — Assessment data (JSON) for modals
# ─────────────────────────────────────────────

@admin_bp.route('/api/assessments/<int:assessment_id>')
@login_required
def api_assessment(assessment_id):
    assessment = Assessment.query.get_or_404(assessment_id)
    return jsonify(assessment.to_dict())


@admin_bp.route('/api/questions/<int:question_id>')
@login_required
def api_question(question_id):
    q = Question.query.get_or_404(question_id)
    return jsonify(q.to_dict(include_answer=True))


# ─────────────────────────────────────────────
# DAY-BY-DAY ASSESSMENT REPORTS
# ─────────────────────────────────────────────

@admin_bp.route('/reports')
@admin_bp.route('/reports/daily')
@login_required
def daily_reports():
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    assessment_id = request.args.get('assessment_id', type=int)

    reports = _get_daily_report_data(start_date, end_date, assessment_id)

    total_candidates = sum(r['total'] for r in reports)
    total_passed = sum(r['passed'] for r in reports)
    total_failed = sum(r['failed'] for r in reports)
    overall_pass_rate = round((total_passed / total_candidates * 100), 1) if total_candidates > 0 else 0.0

    assessments_list = Assessment.query.order_by(Assessment.title).all()

    return render_template(
        'admin/daily_reports.html',
        reports=reports,
        total_days=len(reports),
        total_candidates=total_candidates,
        total_passed=total_passed,
        total_failed=total_failed,
        overall_pass_rate=overall_pass_rate,
        assessments=assessments_list,
        selected_assessment_id=assessment_id,
        start_date=start_date,
        end_date=end_date
    )


@admin_bp.route('/reports/daily/export')
@login_required
def export_daily_reports():
    fmt = request.args.get('fmt', 'xlsx').lower()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    assessment_id = request.args.get('assessment_id', type=int)

    if fmt == 'csv':
        return export_daily_reports_csv(start_date=start_date, end_date=end_date, assessment_id=assessment_id)
    return export_daily_reports_xlsx(start_date=start_date, end_date=end_date, assessment_id=assessment_id)


@admin_bp.route('/reports/tech-round-2')
@login_required
def report_tech_round_2():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '').strip()
    status_filter = request.args.get('status', 'all').strip().lower()
    track_filter = request.args.get('track', type=int)
    per_page = 25

    query = (
        db.session.query(Submission)
        .options(
            joinedload(Submission.candidate),
            joinedload(Submission.assessment)
        )
        .join(Assessment, Submission.assessment_id == Assessment.id)
        .filter(
            Submission.status != 'in_progress',
            db.or_(
                Assessment.title.ilike('%Round 2%'),
                Assessment.title.ilike('%Technical Round 2%')
            ),
            ~Assessment.title.ilike('%Round 3%'),
            ~Assessment.title.ilike('%Coding%')
        )
    )

    if status_filter in ('pass', 'fail'):
        query = query.filter(Submission.status == status_filter)

    if track_filter:
        query = query.filter(Submission.assessment_id == track_filter)

    if search:
        like = f'%{search}%'
        query = query.join(Submission.candidate).filter(
            db.or_(
                Candidate.full_name.ilike(like),
                Candidate.hall_ticket.ilike(like),
                Candidate.email.ilike(like),
            )
        )

    query = query.order_by(
        Submission.percentage.desc(),
        Submission.score.desc(),
        Submission.submitted_at.desc()
    )

    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    round2_assessments = (
        Assessment.query
        .filter(
            db.or_(
                Assessment.title.ilike('%Round 2%'),
                Assessment.title.ilike('%Technical Round 2%')
            ),
            ~Assessment.title.ilike('%Round 3%'),
            ~Assessment.title.ilike('%Coding%')
        )
        .order_by(Assessment.id)
        .all()
    )

    return render_template(
        'admin/reports_tech_2.html',
        pagination=pagination,
        reports=pagination.items,
        round2_assessments=round2_assessments,
        selected_track=track_filter,
        search=search,
        status=status_filter
    )


@admin_bp.route('/reports/tech-round-2/export/<format_type>')
@login_required
def export_tech_round_2_reports(format_type):
    from io import BytesIO, StringIO
    import csv

    status_filter = request.args.get('status', '').strip().lower()
    track_filter = request.args.get('track', type=int)

    query_base = (
        db.session.query(Submission)
        .options(
            joinedload(Submission.candidate),
            joinedload(Submission.assessment)
        )
        .join(Assessment, Submission.assessment_id == Assessment.id)
        .filter(
            Submission.status != 'in_progress',
            db.or_(
                Assessment.title.ilike('%Round 2%'),
                Assessment.title.ilike('%Technical Round 2%')
            ),
            ~Assessment.title.ilike('%Round 3%'),
            ~Assessment.title.ilike('%Coding%')
        )
    )

    if status_filter in ('pass', 'passed'):
        query_base = query_base.filter(Submission.status == 'pass')
        file_suffix = "PASSED_Candidates"
    elif status_filter in ('fail', 'failed'):
        query_base = query_base.filter(Submission.status == 'fail')
        file_suffix = "FAILED_Candidates"
    else:
        file_suffix = "All_Candidates"

    if track_filter:
        query_base = query_base.filter(Submission.assessment_id == track_filter)

    query = query_base.order_by(Submission.submitted_at.desc()).all()

    if format_type == 'csv':
        si = StringIO()
        cw = csv.writer(si)
        cw.writerow(['Candidate Name', 'Email', 'Hall Ticket', 'Track / Assessment', 'Score', 'Total Questions', 'Percentage', 'Status', 'Violations', 'Submitted At (IST)'])
        for r in query:
            cw.writerow([
                r.candidate.full_name if r.candidate else 'N/A',
                r.candidate.email if r.candidate else 'N/A',
                r.candidate.hall_ticket if r.candidate else 'N/A',
                r.assessment.title if r.assessment else 'Round 2 FIB',
                r.score,
                r.total_questions,
                f"{r.percentage:.1f}%",
                r.status.upper(),
                r.violations,
                r.submitted_at_ist.strftime('%Y-%m-%d %I:%M:%S %p') if r.submitted_at_ist else 'N/A'
            ])
        output = make_response(si.getvalue())
        output.headers["Content-Disposition"] = f"attachment; filename=Technical_Round_2_{file_suffix}.csv"
        output.headers["Content-type"] = "text/csv"
        return output

    elif format_type in ('xlsx', 'excel'):
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Tech Round 2 FIB Results"
            ws.append(['Candidate Name', 'Email', 'Hall Ticket', 'Track / Assessment', 'Score', 'Total Questions', 'Percentage', 'Status', 'Violations', 'Submitted At (IST)'])
            for r in query:
                ws.append([
                    r.candidate.full_name if r.candidate else 'N/A',
                    r.candidate.email if r.candidate else 'N/A',
                    r.candidate.hall_ticket if r.candidate else 'N/A',
                    r.assessment.title if r.assessment else 'Round 2 FIB',
                    r.score,
                    r.total_questions,
                    f"{r.percentage:.1f}%",
                    r.status.upper(),
                    r.violations,
                    r.submitted_at_ist.strftime('%Y-%m-%d %I:%M:%S %p') if r.submitted_at_ist else 'N/A'
                ])
            out = BytesIO()
            wb.save(out)
            out.seek(0)
            output = make_response(out.getvalue())
            output.headers["Content-Disposition"] = f"attachment; filename=Technical_Round_2_{file_suffix}.xlsx"
            output.headers["Content-type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            return output
        except Exception:
            si = StringIO()
            cw = csv.writer(si)
            cw.writerow(['Candidate Name', 'Email', 'Hall Ticket', 'Track / Assessment', 'Score', 'Total Questions', 'Percentage', 'Status', 'Violations', 'Submitted At (IST)'])
            for r in query:
                cw.writerow([
                    r.candidate.full_name if r.candidate else 'N/A',
                    r.candidate.email if r.candidate else 'N/A',
                    r.candidate.hall_ticket if r.candidate else 'N/A',
                    r.assessment.title if r.assessment else 'Round 2 FIB',
                    r.score,
                    r.total_questions,
                    f"{r.percentage:.1f}%",
                    r.status.upper(),
                    r.violations,
                    r.submitted_at_ist.strftime('%Y-%m-%d %I:%M:%S %p') if r.submitted_at_ist else 'N/A'
                ])
            output = make_response(si.getvalue())
            output.headers["Content-Disposition"] = f"attachment; filename=Technical_Round_2_{file_suffix}.csv"
            output.headers["Content-type"] = "text/csv"
            return output

    return redirect(url_for('admin.report_tech_round_2'))


@admin_bp.route('/api/submissions/<int:submission_id>/fib-details')
@login_required
def api_submission_fib_details(submission_id):
    submission = Submission.query.get_or_404(submission_id)
    questions = Question.query.filter_by(assessment_id=submission.assessment_id).order_by(Question.id).all()
    answers = {a.question_id: a.selected_option for a in submission.answers}

    detail_list = []
    for idx, q in enumerate(questions, 1):
        user_ans = (answers.get(q.id) or '').strip()
        expected = (q.correct_answer or '').strip()
        expected_list = [ans.strip().lower() for ans in expected.split('|')]
        if len(expected_list) == 1 and ',' in expected:
            expected_list = [ans.strip().lower() for ans in expected.split(',')]

        is_correct = (user_ans.lower() in expected_list) if user_ans else False
        detail_list.append({
            'num': idx,
            'question': q.question,
            'user_answer': user_ans or '[BLANK / UNANSWERED]',
            'expected_answer': expected,
            'is_correct': is_correct
        })

    return jsonify({
        'submission_id': submission.id,
        'candidate_name': submission.candidate.full_name if submission.candidate else 'Unknown',
        'hall_ticket': submission.candidate.hall_ticket if submission.candidate else 'N/A',
        'assessment_title': submission.assessment.title if submission.assessment else 'Technical Round 2',
        'score': submission.score,
        'total': submission.total_questions,
        'percentage': submission.percentage,
        'status': submission.status,
        'violations': submission.violations,
        'questions': detail_list
    })


# ─────────────────────────────────────────────
# TECHNICAL ROUND 3 (LIVE CODING) RESULTS
# ─────────────────────────────────────────────

@admin_bp.route('/reports/tech-round-3')
@login_required
def report_tech_round_3():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '').strip()
    status_filter = request.args.get('status', 'all').strip().lower()
    track_filter = request.args.get('track', type=int)
    per_page = 25

    query = (
        db.session.query(Submission)
        .options(
            joinedload(Submission.candidate),
            joinedload(Submission.assessment),
            joinedload(Submission.coding_submissions).joinedload(CodingSubmission.problem)
        )
        .join(Assessment, Submission.assessment_id == Assessment.id)
        .filter(
            Submission.status != 'in_progress',
            db.or_(
                Assessment.title.ilike('%Round 3%'),
                Assessment.title.ilike('%Coding%'),
                Assessment.id.in_([4, 20, 21, 22, 23, 24, 25])
            )
        )
    )

    if status_filter in ('pass', 'fail'):
        query = query.filter(Submission.status == status_filter)

    if track_filter:
        query = query.filter(Submission.assessment_id == track_filter)

    if search:
        like = f'%{search}%'
        query = query.join(Submission.candidate).filter(
            db.or_(
                Candidate.full_name.ilike(like),
                Candidate.hall_ticket.ilike(like),
                Candidate.email.ilike(like),
            )
        )

    query = query.order_by(
        Submission.percentage.desc(),
        Submission.score.desc(),
        Submission.submitted_at.desc()
    )

    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    round3_assessments = (
        Assessment.query
        .filter(
            db.or_(
                Assessment.title.ilike('%Round 3%'),
                Assessment.title.ilike('%Coding%'),
                Assessment.id.in_([4, 20, 21, 22, 23, 24, 25])
            )
        )
        .order_by(Assessment.id)
        .all()
    )

    # Compute aggregate stats for Round 3
    all_r3_subs = (
        db.session.query(Submission)
        .join(Assessment, Submission.assessment_id == Assessment.id)
        .filter(
            Submission.status != 'in_progress',
            db.or_(
                Assessment.title.ilike('%Round 3%'),
                Assessment.title.ilike('%Coding%'),
                Assessment.id.in_([4, 20, 21, 22, 23, 24, 25])
            )
        )
        .all()
    )
    total_count = len(all_r3_subs)
    passed_count = sum(1 for s in all_r3_subs if s.status == 'pass')
    failed_count = sum(1 for s in all_r3_subs if s.status == 'fail')
    pass_rate = round((passed_count / total_count * 100), 1) if total_count > 0 else 0.0
    avg_score = round(sum(s.score or s.coding_score or 0 for s in all_r3_subs) / total_count, 1) if total_count > 0 else 0.0

    return render_template(
        'admin/reports_tech_3.html',
        pagination=pagination,
        reports=pagination.items,
        round3_assessments=round3_assessments,
        selected_track=track_filter,
        search=search,
        status=status_filter,
        total_count=total_count,
        passed_count=passed_count,
        failed_count=failed_count,
        pass_rate=pass_rate,
        avg_score=avg_score
    )


@admin_bp.route('/reports/tech-round-3/export/<format_type>')
@login_required
def export_tech_round_3_reports(format_type):
    from io import BytesIO, StringIO
    import csv

    status_filter = request.args.get('status', '').strip().lower()
    track_filter = request.args.get('track', type=int)

    query_base = (
        db.session.query(Submission)
        .options(
            joinedload(Submission.candidate),
            joinedload(Submission.assessment),
            joinedload(Submission.coding_submissions)
        )
        .join(Assessment, Submission.assessment_id == Assessment.id)
        .filter(
            Submission.status != 'in_progress',
            db.or_(
                Assessment.title.ilike('%Round 3%'),
                Assessment.title.ilike('%Coding%'),
                Assessment.id.in_([4, 20, 21, 22, 23, 24, 25])
            )
        )
    )

    if status_filter in ('pass', 'passed'):
        query_base = query_base.filter(Submission.status == 'pass')
        file_suffix = "PASSED_Candidates"
    elif status_filter in ('fail', 'failed'):
        query_base = query_base.filter(Submission.status == 'fail')
        file_suffix = "FAILED_Candidates"
    else:
        file_suffix = "All_Candidates"

    if track_filter:
        query_base = query_base.filter(Submission.assessment_id == track_filter)

    query = query_base.order_by(Submission.submitted_at.desc()).all()

    if format_type == 'csv':
        si = StringIO()
        cw = csv.writer(si)
        cw.writerow(['Candidate Name', 'Email', 'Hall Ticket', 'Track / Assessment', 'Coding Score', 'Problems Attempted', 'Percentage', 'Status', 'Violations', 'Submitted At (IST)'])
        for r in query:
            cw.writerow([
                r.candidate.full_name if r.candidate else 'N/A',
                r.candidate.email if r.candidate else 'N/A',
                r.candidate.hall_ticket if r.candidate else 'N/A',
                r.assessment.title if r.assessment else 'Round 3 Coding',
                r.score or r.coding_score or 0,
                len(r.coding_submissions) if r.coding_submissions else r.total_questions,
                f"{r.percentage:.1f}%",
                r.status.upper(),
                r.violations,
                r.submitted_at_ist.strftime('%Y-%m-%d %I:%M:%S %p') if r.submitted_at_ist else 'N/A'
            ])
        output = make_response(si.getvalue())
        output.headers["Content-Disposition"] = f"attachment; filename=Technical_Round_3_{file_suffix}.csv"
        output.headers["Content-type"] = "text/csv"
        return output

    elif format_type in ('xlsx', 'excel'):
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Tech Round 3 Coding Results"
            ws.append(['Candidate Name', 'Email', 'Hall Ticket', 'Track / Assessment', 'Coding Score', 'Problems Attempted', 'Percentage', 'Status', 'Violations', 'Submitted At (IST)'])
            for r in query:
                ws.append([
                    r.candidate.full_name if r.candidate else 'N/A',
                    r.candidate.email if r.candidate else 'N/A',
                    r.candidate.hall_ticket if r.candidate else 'N/A',
                    r.assessment.title if r.assessment else 'Round 3 Coding',
                    r.score or r.coding_score or 0,
                    len(r.coding_submissions) if r.coding_submissions else r.total_questions,
                    f"{r.percentage:.1f}%",
                    r.status.upper(),
                    r.violations,
                    r.submitted_at_ist.strftime('%Y-%m-%d %I:%M:%S %p') if r.submitted_at_ist else 'N/A'
                ])
            out = BytesIO()
            wb.save(out)
            out.seek(0)
            output = make_response(out.getvalue())
            output.headers["Content-Disposition"] = f"attachment; filename=Technical_Round_3_{file_suffix}.xlsx"
            output.headers["Content-type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            return output
        except Exception:
            si = StringIO()
            cw = csv.writer(si)
            cw.writerow(['Candidate Name', 'Email', 'Hall Ticket', 'Track / Assessment', 'Coding Score', 'Problems Attempted', 'Percentage', 'Status', 'Violations', 'Submitted At (IST)'])
            for r in query:
                cw.writerow([
                    r.candidate.full_name if r.candidate else 'N/A',
                    r.candidate.email if r.candidate else 'N/A',
                    r.candidate.hall_ticket if r.candidate else 'N/A',
                    r.assessment.title if r.assessment else 'Round 3 Coding',
                    r.score or r.coding_score or 0,
                    len(r.coding_submissions) if r.coding_submissions else r.total_questions,
                    f"{r.percentage:.1f}%",
                    r.status.upper(),
                    r.violations,
                    r.submitted_at_ist.strftime('%Y-%m-%d %I:%M:%S %p') if r.submitted_at_ist else 'N/A'
                ])
            output = make_response(si.getvalue())
            output.headers["Content-Disposition"] = f"attachment; filename=Technical_Round_3_{file_suffix}.csv"
            output.headers["Content-type"] = "text/csv"
            return output

    return redirect(url_for('admin.report_tech_round_3'))


@admin_bp.route('/api/submissions/<int:submission_id>/coding-details')
@login_required
def api_submission_coding_details(submission_id):
    submission = (
        db.session.query(Submission)
        .options(
            joinedload(Submission.candidate),
            joinedload(Submission.assessment),
            joinedload(Submission.coding_submissions).joinedload(CodingSubmission.problem)
        )
        .filter(Submission.id == submission_id)
        .first_or_404()
    )

    coding_subs = submission.coding_submissions
    problems_list = []

    for cs in coding_subs:
        prob = cs.problem
        problems_list.append({
            'problem_id': cs.problem_id,
            'title': prob.title if prob else 'Coding Problem',
            'difficulty': prob.difficulty if prob else 'Medium',
            'points': prob.points if prob else 100,
            'score': cs.score,
            'language': cs.language,
            'status': cs.status,
            'passed_testcases': cs.passed_testcases,
            'total_testcases': cs.total_testcases,
            'execution_time_ms': cs.execution_time_ms,
            'source_code': cs.source_code,
            'submitted_at': cs.submitted_at.strftime('%d %b %Y, %I:%M %p') if cs.submitted_at else 'N/A'
        })

    # If no coding_submissions exist yet, show problem placeholders from drive
    if not problems_list and submission.assessment_id:
        problems = CodingProblem.query.filter_by(assessment_id=submission.assessment_id).all()
        for p in problems:
            problems_list.append({
                'problem_id': p.id,
                'title': p.title,
                'difficulty': p.difficulty,
                'points': p.points,
                'score': 0,
                'language': 'N/A',
                'status': 'Not Attempted',
                'passed_testcases': 0,
                'total_testcases': len(p.testcases),
                'execution_time_ms': 0,
                'source_code': '// No code submitted for this challenge.',
                'submitted_at': 'N/A'
            })

    return jsonify({
        'submission_id': submission.id,
        'candidate_name': submission.candidate.full_name if submission.candidate else 'Unknown',
        'hall_ticket': submission.candidate.hall_ticket if submission.candidate else 'N/A',
        'email': submission.candidate.email if submission.candidate else 'N/A',
        'assessment_title': submission.assessment.title if submission.assessment else 'Technical Round 3 (Coding)',
        'score': submission.score or submission.coding_score or 0,
        'percentage': submission.percentage,
        'status': submission.status,
        'violations': submission.violations,
        'submitted_at': submission.submitted_at_ist.strftime('%d %b %Y, %I:%M %p') if submission.submitted_at_ist else 'N/A',
        'problems': problems_list
    })
